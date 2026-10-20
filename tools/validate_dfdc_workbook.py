"""Read-only XLSX validation; openpyxl never authors this workbook."""
import json
import math
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]


def main():
    source=json.loads((ROOT/'../research/analyses/dfdc_literature.json').read_text(encoding='utf-8'))
    path=ROOT/'../research/workbooks/DFDC_Research_Comparison.xlsx'
    w=openpyxl.load_workbook(path,data_only=False)
    cached=openpyxl.load_workbook(path,data_only=True)
    s=w['DFDC Published Results']
    assert [c.value for c in s[1]]==source['headers']
    assert s.max_row==len(source['rows'])+1 and s.max_column==20
    assert s.freeze_panes=='C2' and len(s.tables)==1 and not s.merged_cells.ranges
    assert next(iter(s.tables.values())).autoFilter is not None
    for r,item in enumerate(source['rows'],2):
        for col,limit in [(12,100),(13,1),(14,1)]:
            cell=s.cell(r,col);value=cell.value
            assert value is None or (cell.data_type=='n' and isinstance(value,(int,float)) and math.isfinite(value) and 0<=value<=limit)
            assert value==item[source['headers'][col-1]]
        formula=s.cell(r,16).value
        assert formula.startswith('=HYPERLINK(') and item['_key'] in formula
        assert item['_source_path'].split('!')[0].replace('\\','/') in formula
        assert cached[s.title].cell(r,16).value not in {'#REF!','#NAME?','#VALUE!','#N/A'}
        assert all(s.cell(r,c).alignment.wrap_text for c in range(1,21))
    assert w['Source Audit'].max_row==113
    assert w['Local DFDC Baseline']['I1'].value=='Accuracy (%)'
    # Check the two user-identified DFDC results independently of row order.
    assert any(x['Paper ID']=='2' and x['Accuracy (%)']==92.2 and x['ROC-AUC']==.9752 for x in source['rows'])
    assert any(x['Paper ID']=='35' and x['Accuracy (%)']==79 and x['ROC-AUC']==.82 for x in source['rows'])
    for sheet in w:
        assert not sheet.merged_cells.ranges
        for row in sheet:
            for cell in row:
                assert not (isinstance(cell.value,str) and '\ufffd' in cell.value)
    report={'passed':True,'published_rows':len(source['rows']),'experiment_families':len({x['_key'] for x in source['rows']}),
            'audited_unique_pdfs':len(source['audit']),'checks':'Exact ordered headers; numeric bounds/values; source hyperlinks; filter table; freeze panes; wrapping; no merged cells; known DFDC scores; Unicode; cached formula errors'}
    (ROOT/'../research/analyses/workbook_validation.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report,indent=2))


if __name__=='__main__':main()
