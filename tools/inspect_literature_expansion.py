"""Read-only structural and content inventory of the authoritative backup."""
import json
import sys
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'research/working/expansion'
OUT.mkdir(parents=True,exist_ok=True)
records=sorted((ROOT/'archive/literature_backups').glob('backup_*.json'))
backup=json.loads(records[-1].read_text(encoding='utf-8-sig'))
SOURCE=Path(backup['Backup'])
w=openpyxl.load_workbook(SOURCE,data_only=False)
wc=openpyxl.load_workbook(SOURCE,data_only=True)
result={'source':str(SOURCE),'backup_record':backup,'sheets':{},'names':list(w.defined_names)}
for s in w:
    cells={c.coordinate:{'value':c.value,'type':c.data_type,'style':c.style_id,'number_format':c.number_format,
                         'cached':wc[s.title][c.coordinate].value} for row in s for c in row if c.value is not None}
    result['sheets'][s.title]={'rows':s.max_row,'columns':s.max_column,'state':s.sheet_state,'freeze':s.freeze_panes,
                             'merges':list(map(str,s.merged_cells.ranges)),'filter':s.auto_filter.ref,
                             'tables':{v.name:v.ref for v in s.tables.values()} if s.tables else {},
                             'charts':[{'type':c.__class__.__name__,'anchor':str(c.anchor),'series':str(c.ser)} for c in s._charts],
                             'conditional_formatting':str(list(s.conditional_formatting)),'cells':cells,
                             'row_heights':{k:v.height for k,v in s.row_dimensions.items()},
                             'column_widths':{k:v.width for k,v in s.column_dimensions.items()}}
    print(s.title, s.max_row,s.max_column,'formulas',sum(c['type']=='f' for c in cells.values()),'charts',len(s._charts),'merges',len(s.merged_cells.ranges),'filter',s.auto_filter.ref)
with zipfile.ZipFile(SOURCE) as z:
    result['package_parts']={i.filename:{'bytes':i.file_size,'crc':i.CRC} for i in z.infolist()}
(OUT/'original_inventory.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
master=w['Master Literature Database']
headers=[c.value for c in master[1]]
rows=[dict(zip(headers,row)) for row in master.iter_rows(min_row=2,values_only=True) if isinstance(row[0],(int,float))]
(OUT/'original_master.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
print('Master records',len(rows),'types',dict(Counter(r['Paper Type'] for r in rows)))
print('Parts',len(result['package_parts']),'special', [p for p in result['package_parts'] if any(k in p for k in ['chart','drawing','pivot','vml','comment','external'])])
if len(sys.argv)>1:
    for name in sys.argv[1:]:
        for coord,c in result['sheets'][name]['cells'].items():print(coord,str(c['value'])[:700])
