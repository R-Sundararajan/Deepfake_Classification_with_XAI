"""Read-only validation of the two final six-paper deliverables."""
import hashlib
import json
import math
import re
from pathlib import Path
import openpyxl

PROJECT=Path(__file__).resolve().parents[1]
WORK=PROJECT/'../research/working/existing_six'
data=json.loads((WORK/'analysis.json').read_text(encoding='utf-8'))
native=json.loads((WORK/'native_validation.json').read_text(encoding='utf-8-sig'))
output=PROJECT/'../research/workbooks/Existing_Six_Papers_Detailed_Analysis.xlsx'
report=PROJECT/'../research/reports/Existing_Six_Papers_Deep_Analysis.md'
wf=openpyxl.load_workbook(output,data_only=False)
wv=openpyxl.load_workbook(output,data_only=True)
checks=[]
def test(label,condition):
    checks.append({'check':label,'passed':bool(condition)})

test('Exactly nine requested sheets',wf.sheetnames==['Summary','DFDC Verified','Hardware Feasibility','2','9','15','16','35','45'])
test('Exactly six summary records and stable requested IDs',[wv['Summary'].cell(r,1).value for r in range(6,12)]==[2,9,15,16,35,45])
test('Complete original PDF reading coverage',sum(p['pages'] for p in data['papers'])==63)
test('Original literature workbook unchanged',hashlib.sha256(Path(data['source_workbook']).read_bytes()).hexdigest().upper()==data['source_sha256'])
md=report.read_text(encoding='utf-8')
test('Exactly six paper sections in report',list(map(int,re.findall(r'^## Paper (\d+):',md,re.M)))==[2,9,15,16,35,45])
for p in data['papers']:
    pid=p['paper_id'];s=wf[str(pid)]
    test(f'Paper {pid}: source PDF hash matches',hashlib.sha256(Path(p['source']).read_bytes()).hexdigest()==p['sha256'])
    test(f'Paper {pid}: traceable source in sheet',p['source'] in s['A3'].value and p['sha256'] in s['A4'].value)
    test(f'Paper {pid}: report provenance',p['sha256'] in md and Path(p['source']).name in md)
    test(f'Paper {pid}: filter and frozen headers',bool(s.tables) and s.freeze_panes=='C7')
    test(f'Paper {pid}: header row visible',not s.row_dimensions[6].hidden and s['A6'].value=='Topic')
    test(f'Paper {pid}: all source-grounded notes present',all(n['text'] in md for n in p['notes']))
    test(f'Paper {pid}: required eight analysis areas',{'Architecture','Explainability','Dataset','Performance','Compute','Optimization','Limitations','Future relevance'} <= {n['section'] for n in p['notes']} | ({'Performance'} if p['performance'] else set()))
    test(f'Paper {pid}: citations on every detail row',all(s.cell(r,7).value for r in range(7,s.max_row+1)))
    test(f'Paper {pid}: evidence status on every detail row',all(s.cell(r,6).value in ['Reported','Reported (review)','Calculated','Estimated','Unknown','Interpretation'] for r in range(7,s.max_row+1)))

verified=data['dfdc_performance']
test('DFDC comparison contains only two experimental papers',{wv['DFDC Verified'].cell(r,1).value for r in range(6,14)}=={2,35} and len(verified)==8)
for row,r in enumerate(verified,6):
    for col,key in enumerate(['accuracy','precision','recall','f1','auc'],4):
        v=wv['DFDC Verified'].cell(row,col).value
        test(f'DFDC row {row} {key}',v=='Not Reported' if r[key] is None else isinstance(v,(int,float)) and math.isclose(v,r[key],abs_tol=1e-12))
    test(f'DFDC row {row}: numeric headline and source',wv['DFDC Verified'].cell(row,4).data_type=='n' and wv['DFDC Verified'].cell(row,10).value==r['reference'])
    nr=native['DFDC'][row-6]
    test(f'DFDC row {row}: native Excel agreement',math.isclose(nr['Accuracy'],r['accuracy'],abs_tol=1e-12) and math.isclose(nr['AUC'],r['auc'],abs_tol=1e-12))
for cell,expected in {'C22':262144,'C23':24.4140625,'C26':15.2587890625,'C30':5080,'C36':304800000,'C38':846.6666666666666,'C43':80/2.17,'C44':60/2.17}.items():
    test(f'Hardware formula {cell}',math.isclose(wv['Hardware Feasibility'][cell].value,expected,abs_tol=1e-10))
test('Native Excel calculation and scenario change',not native['FormulaErrors'] and native['ReadOnly'] and not native['Saved'] and math.isclose(native['NativeBaseHours'],2*native['NativeChangedHours'],abs_tol=1e-10) and native['ScenarioRestored']==100)
errors=[]
for s in wf:
    for row in s:
        for c in row:
            if c.data_type=='e' or (c.data_type=='f' and '#REF!' in c.value):errors.append(f'{s.title}!{c.coordinate}')
    test(f'{s.title}: row heights within Excel limit',all((d.height or 15)<=409.5 for d in s.row_dimensions.values()))
test('No saved formula errors or broken references',not errors)
result={'status':'PASS' if all(x['passed'] for x in checks) else 'FAIL','checks':checks,'check_count':len(checks),
        'original_workbook_unchanged':True,'original_sha256':data['source_sha256'],
        'output_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'report_sha256':hashlib.sha256(report.read_bytes()).hexdigest(),
        'native_excel_version':native['ExcelVersion'],'sheets':wf.sheetnames,'pdf_pages_read':63,
        'scientific_limit':'Source-verified reported claims and explicit interpretations; no independent reproduction, code review or measured RTX 4060 benchmark.'}
(WORK/'final_validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps({k:result[k] for k in ['status','check_count','original_workbook_unchanged','sheets','pdf_pages_read']},indent=2))
for x in checks:
    if not x['passed']:print('FAILED',x['check'])
raise SystemExit(0 if result['status']=='PASS' else 1)
